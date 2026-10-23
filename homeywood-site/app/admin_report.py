from datetime import datetime, timedelta

from flask import Blueprint, jsonify, request, send_file
from sqlalchemy import func
from sqlalchemy.orm import joinedload

from .catalog import admin_required
from .extensions import db
from .models import Category, Order, OrderItem, Product, User
from .orders import STATUS_LABELS

admin_report_bp = Blueprint("admin_report", __name__, url_prefix="/api")

REVENUE_STATUSES = ("diproses", "dikirim", "selesai")
CAT_ORDER = ["Sofa", "Meja", "Kursi", "Lemari"]


def _parse_range():
    """Ambil ?from=YYYY-MM-DD&to=YYYY-MM-DD. Default: awal bulan ini s/d hari ini."""
    today = datetime.now().date()
    try:
        d_from = datetime.strptime(request.args.get("from", ""), "%Y-%m-%d").date()
    except ValueError:
        d_from = today.replace(day=1)
    try:
        d_to = datetime.strptime(request.args.get("to", ""), "%Y-%m-%d").date()
    except ValueError:
        d_to = today
    if d_to < d_from:
        d_from, d_to = d_to, d_from
    if (d_to - d_from).days > 366:
        d_from = d_to - timedelta(days=366)
    return d_from, d_to


def _sort_cats(rows):
    return sorted(rows, key=lambda r: (CAT_ORDER.index(r["name"]) if r["name"] in CAT_ORDER else 99, r["name"]))


def build_report(d_from, d_to):
    start = datetime.combine(d_from, datetime.min.time())
    end = datetime.combine(d_to + timedelta(days=1), datetime.min.time())  # eksklusif

    in_range = db.and_(
        Order.created_at >= start, Order.created_at < end,
        Order.status.in_(REVENUE_STATUSES),
    )

    total_revenue = float(
        db.session.query(func.coalesce(func.sum(Order.total_amount), 0)).filter(in_range).scalar() or 0
    )
    total_orders = Order.query.filter(in_range).count()

    cat_name = func.coalesce(Category.name, "Lainnya")
    rows = (
        db.session.query(
            cat_name.label("name"),
            func.coalesce(func.sum(OrderItem.quantity), 0),
            func.count(func.distinct(OrderItem.order_id)),
            func.coalesce(func.sum(OrderItem.subtotal), 0),
        )
        .join(Order, Order.id == OrderItem.order_id)
        .outerjoin(Product, Product.id == OrderItem.product_id)
        .outerjoin(Category, Category.id == Product.category_id)
        .filter(in_range)
        .group_by(cat_name)
        .all()
    )
    sold = {r[0]: {"name": r[0], "qty": int(r[1]), "orders": int(r[2]), "revenue": float(r[3])} for r in rows}
    recap = [sold.get(c, {"name": c, "qty": 0, "orders": 0, "revenue": 0.0}) for c in CAT_ORDER]
    recap += _sort_cats([v for k, v in sold.items() if k not in CAT_ORDER])

    prod_rows = (
        db.session.query(func.coalesce(Category.name, "Lainnya"), func.count(Product.id))
        .select_from(Product)
        .outerjoin(Category, Category.id == Product.category_id)
        .filter(Product.is_active.is_(True))
        .group_by(func.coalesce(Category.name, "Lainnya"))
        .all()
    )
    prod_map = {r[0]: int(r[1]) for r in prod_rows}
    prod_by_cat = [{"name": c, "value": prod_map.get(c, 0)} for c in CAT_ORDER]
    prod_by_cat += _sort_cats([{"name": k, "value": v} for k, v in prod_map.items() if k not in CAT_ORDER])

    active = User.query.filter_by(role="customer", is_active=True).count()
    inactive = User.query.filter_by(role="customer", is_active=False).count()

    ship_total = total_revenue - sum(r["revenue"] for r in recap)
    recap_table = recap + (
        [{"name": "Ongkos kirim", "qty": 0, "orders": 0, "revenue": ship_total}]
        if ship_total > 0.5 else []
    )

    return {
        "from": d_from.isoformat(),
        "to": d_to.isoformat(),
        "revenue": {"total": total_revenue, "by_cat": [{"name": r["name"], "value": r["revenue"]} for r in recap]},
        "orders": {"total": total_orders, "by_cat": [{"name": r["name"], "value": r["orders"]} for r in recap]},
        "products": {"total": sum(p["value"] for p in prod_by_cat), "by_cat": prod_by_cat},
        "users": {"total": active + inactive, "active": active, "inactive": inactive},
        "recap": recap_table,
        "_range": (start, end),
    }


def _order_rows(start, end):
    orders = (
        Order.query.options(joinedload(Order.user)).filter(Order.created_at >= start, Order.created_at < end, Order.status.in_(REVENUE_STATUSES))
        .order_by(Order.created_at.asc(), Order.id.asc())
        .all()
    )
    return [
        [
            o.order_code,
            o.created_at.strftime("%d-%m-%Y %H:%M") if o.created_at else "",
            o.user.full_name if o.user else "-",
            STATUS_LABELS.get(o.status, o.status),
            float(o.total_amount),
        ]
        for o in orders
    ]


def _rp(n):
    return "Rp " + f"{int(round(n)):,}".replace(",", ".")

def _xl(value):
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value

@admin_report_bp.route("/admin/report", methods=["GET"])
@admin_required
def admin_report():
    d_from, d_to = _parse_range()
    data = build_report(d_from, d_to)
    data.pop("_range")
    return jsonify(data)


@admin_report_bp.route("/admin/report/export.xlsx", methods=["GET"])
@admin_required
def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    d_from, d_to = _parse_range()
    data = build_report(d_from, d_to)
    start, end = data["_range"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap"
    head_fill = PatternFill("solid", fgColor="EFE4D2")
    bold = Font(bold=True)

    ws["A1"] = "Laporan Penjualan Homey Wood"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Periode: {d_from.strftime('%d-%m-%Y')} s/d {d_to.strftime('%d-%m-%Y')}"
    ws["A4"], ws["B4"] = "Total Pendapatan", data["revenue"]["total"]
    ws["A5"], ws["B5"] = "Total Transaksi", data["orders"]["total"]
    ws["A6"], ws["B6"] = "Produk Terdaftar", data["products"]["total"]
    ws["A7"], ws["B7"] = "Total Pengguna", data["users"]["total"]
    for r in range(4, 8):
        ws[f"A{r}"].font = bold
    ws["B4"].number_format = '"Rp" #,##0'

    ws.append([])
    ws.append(["Kategori", "Terjual (unit)", "Transaksi", "Pendapatan"])
    hdr = ws.max_row
    for c in ws[hdr]:
        c.font, c.fill = bold, head_fill
    for r in data["recap"]:
        ws.append([r["name"], r["qty"], r["orders"], r["revenue"]])
        ws.cell(row=ws.max_row, column=4).number_format = '"Rp" #,##0'
    ws.append(["Total", sum(r["qty"] for r in data["recap"]), data["orders"]["total"], data["revenue"]["total"]])
    for c in ws[ws.max_row]:
        c.font = bold
    ws.cell(row=ws.max_row, column=4).number_format = '"Rp" #,##0'
    for col, w in zip("ABCD", (24, 16, 14, 20)):
        ws.column_dimensions[col].width = w
    ws["B4"].alignment = Alignment(horizontal="left")

    ws2 = wb.create_sheet("Transaksi")
    ws2.append(["ID Pesanan", "Tanggal", "Pelanggan", "Status", "Total"])
    for c in ws2[1]:
        c.font, c.fill = bold, head_fill
    for row in _order_rows(start, end):
        row[2] = _xl(row[2])
        ws2.append(row)
        ws2.cell(row=ws2.max_row, column=5).number_format = '"Rp" #,##0'
    for col, w in zip("ABCDE", (22, 18, 26, 16, 18)):
        ws2.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"laporan-penjualan_{d_from.isoformat()}_{d_to.isoformat()}.xlsx"
    return send_file(buf, as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@admin_report_bp.route("/admin/report/export.pdf", methods=["GET"])
@admin_required
def export_pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    d_from, d_to = _parse_range()
    data = build_report(d_from, d_to)
    start, end = data["_range"]
    styles = getSampleStyleSheet()

    def table(rows, widths, total_row=False):
        t = Table(rows, colWidths=widths, repeatRows=1)
        st = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EFE4D2")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CDBFA8")),
            ("ALIGN", (-1, 0), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]
        if total_row:
            st += [("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                   ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F6EFE3"))]
        t.setStyle(TableStyle(st))
        return t

    story = [
        Paragraph("Laporan Penjualan Homey Wood", styles["Title"]),
        Paragraph(f"Periode: {d_from.strftime('%d-%m-%Y')} s/d {d_to.strftime('%d-%m-%Y')}", styles["Normal"]),
        Spacer(1, 12),
        table([
            ["Ringkasan", ""],
            ["Total Pendapatan", _rp(data["revenue"]["total"])],
            ["Total Transaksi", str(data["orders"]["total"])],
            ["Produk Terdaftar", str(data["products"]["total"])],
            ["Total Pengguna", str(data["users"]["total"])],
        ], [8 * cm, 8 * cm]),
        Spacer(1, 14),
        Paragraph("Rekap per Kategori", styles["Heading3"]),
        table(
            [["Kategori", "Terjual", "Transaksi", "Pendapatan"]]
            + [[r["name"], r["qty"], r["orders"], _rp(r["revenue"])] for r in data["recap"]]
            + [["Total", sum(r["qty"] for r in data["recap"]), data["orders"]["total"], _rp(data["revenue"]["total"])]],
            [5 * cm, 3 * cm, 3 * cm, 5 * cm], total_row=True),
        Spacer(1, 14),
        Paragraph("Daftar Transaksi", styles["Heading3"]),
    ]
    rows = _order_rows(start, end)
    if rows:
        story.append(table(
            [["ID Pesanan", "Tanggal", "Pelanggan", "Status", "Total"]]
            + [[r[0], r[1], r[2][:22], r[3], _rp(r[4])] for r in rows],
            [3.6 * cm, 3.2 * cm, 4.4 * cm, 2.6 * cm, 3.2 * cm]))
    else:
        story.append(Paragraph("Tidak ada transaksi pada periode ini.", styles["Normal"]))

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                      bottomMargin=2 * cm, title="Laporan Penjualan Homey Wood").build(story)
    buf.seek(0)
    name = f"laporan-penjualan_{d_from.isoformat()}_{d_to.isoformat()}.pdf"
    return send_file(buf, as_attachment=True, download_name=name, mimetype="application/pdf")
